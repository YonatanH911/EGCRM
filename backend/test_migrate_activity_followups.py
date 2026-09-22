import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker

import migrate_activity_followups
from migrate_activity_followups import load_followups, normalize_email
import models


class ActivityFollowupMigrationTests(unittest.TestCase):
    def test_normalizes_email(self):
        self.assertEqual(normalize_email("  Person@Example.COM  "), "person@example.com")

    def test_loads_guids_and_emails_without_guessing_conflicts(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Activities.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["(Do Not Modify) Activity", "Subject", "Follow Up Email"])
            sheet.append(["{ABC}", "First", " Person@Example.COM "])
            sheet.append(["def", "Second", "second@example.com"])
            sheet.append(["DEF", "Second", "different@example.com"])
            sheet.append(["ghi", "Third", None])
            workbook.save(path)
            workbook.close()

            export = load_followups(path)

        self.assertEqual(export.emails, {"abc": "person@example.com"})
        self.assertEqual(export.rows_scanned, 4)
        self.assertEqual(export.blank_email, 1)
        self.assertEqual(export.conflicting_guids, {"def"})

    def test_requires_follow_up_column(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Activities.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.append(["(Do Not Modify) Activity", "Subject"])
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "Follow Up Email"):
                load_followups(path)

    def test_dry_run_and_existing_links_are_preserved(self):
        engine = create_engine("sqlite:///:memory:")
        models.Base.metadata.create_all(engine)
        with engine.begin() as connection:
            connection.execute(text(
                "CREATE TABLE legacy_import_ids (entity_type VARCHAR(32), source_id VARCHAR(64), local_id INT)"
            ))
        session_factory = sessionmaker(bind=engine)
        with session_factory() as db:
            original = models.Contact(first_name="Original", last_name="Person", email="old@example.com")
            replacement = models.Contact(first_name="New", last_name="Person", email="new@example.com")
            activity = models.Activity(subject="Call", contact=original)
            unlinked = models.Activity(subject="Email")
            db.add_all([original, replacement, activity, unlinked])
            db.flush()
            activity_id = activity.id
            unlinked_id = unlinked.id
            original_id = original.id
            replacement_id = replacement.id
            db.execute(text(
                "INSERT INTO legacy_import_ids VALUES ('activity', 'abc', :local_id)"
            ), {"local_id": activity_id})
            db.execute(text(
                "INSERT INTO legacy_import_ids VALUES ('activity', 'def', :local_id)"
            ), {"local_id": unlinked_id})
            db.commit()

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Activities.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.append(["(Do Not Modify) Activity", "Follow Up Email"])
            workbook.active.append(["abc", "new@example.com"])
            workbook.active.append(["def", "new@example.com"])
            workbook.save(path)
            workbook.close()

            with patch.object(migrate_activity_followups, "SessionLocal", session_factory):
                with contextlib.redirect_stdout(io.StringIO()):
                    migrate_activity_followups.run(path)
                with session_factory() as db:
                    self.assertEqual(db.get(models.Activity, activity_id).contact_id, original_id)
                    self.assertIsNone(db.get(models.Activity, unlinked_id).contact_id)

                with contextlib.redirect_stdout(io.StringIO()):
                    migrate_activity_followups.run(path, apply=True)
                with session_factory() as db:
                    self.assertEqual(db.get(models.Activity, activity_id).contact_id, original_id)
                    self.assertEqual(db.get(models.Activity, unlinked_id).contact_id, replacement_id)

                with contextlib.redirect_stdout(io.StringIO()):
                    migrate_activity_followups.run(path, apply=True, overwrite=True)
                with session_factory() as db:
                    self.assertEqual(db.get(models.Activity, activity_id).contact_id, replacement_id)


if __name__ == "__main__":
    unittest.main()
